from pathlib import Path

from backend.models.schemas import AnalysisResponse


def generate_approval_note(result: AnalysisResponse, output_dir: Path) -> Path:
    try:
        from docx import Document

        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / "Approval_Note.docx"
        document = Document()
        document.add_heading("SOVEREIGN-X", 0)
        document.add_heading("ENGINEERING APPROVAL NOTE", 1)
        rows = [
            ("Equipment", result.finding.equipment),
            ("Finding", result.ai_analysis.finding),
            ("Measured Value", f"{result.calculation.measured} mm"),
            ("Applicable Requirement", f"{result.sop_evidence.document} §{result.sop_evidence.section}"),
            ("Permitted Limit", f"{result.calculation.limit} mm"),
            ("Verification", result.calculation.expression),
            ("Verification Status", result.calculation.status),
            ("AI Recommendation", result.ai_analysis.recommendation),
            ("Source Evidence", f"{result.finding.document} — Page {result.finding.page} — {result.finding.photo}"),
            ("AI Model", f"{result.ai_analysis.model} — {result.processing_mode}"),
            ("Confidence", result.ai_analysis.confidence),
            ("Processing Mode", "LOCAL / OFFLINE" if result.processing_mode == "LOCAL" else "DEMO / SYNTHETIC"),
        ]
        for label, value in rows:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(value)
        document.add_paragraph("\nAI-GENERATED DRAFT\nENGINEER REVIEW REQUIRED")
        document.save(path)
        return path
    except Exception as exc:
        raise RuntimeError(f"Approval note generation failed: {exc}") from exc

