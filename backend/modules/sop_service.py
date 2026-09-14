"""Local SOP management, lifecycle, and retrieval service."""
from __future__ import annotations

import hashlib
import re
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from backend.modules.embeddings import embed
from backend.modules.ocr import extract_document
from backend.modules.storage import ObjectStorage
from backend.db import SessionLocal
from backend.models.db_models import AuditLog, Document, SOP, SOPChunk, SOPDocument

SOP_STATUSES = {"UPLOADED", "UNDER_REVIEW", "APPROVED", "ACTIVE", "SUPERSEDED", "ARCHIVED"}
TRANSITIONS = {
    "submit-review": ("UPLOADED", "UNDER_REVIEW"),
    "approve": ("UNDER_REVIEW", "APPROVED"),
    "activate": ("APPROVED", "ACTIVE"),
    "supersede": ("ACTIVE", "SUPERSEDED"),
    "archive": ("SUPERSEDED", "ARCHIVED"),
}
ROLE_ACTIONS = {
    "submit-review": {"ADMIN", "ENGINEER", "SUPERVISOR"},
    "approve": {"ADMIN", "SUPERVISOR"},
    "activate": {"ADMIN", "SUPERVISOR"},
    "supersede": {"ADMIN", "SUPERVISOR"},
    "archive": {"ADMIN", "AUDITOR", "SUPERVISOR"},
}


@dataclass
class SOPChunkRecord:
    id: str
    sop_id: str
    sop_version: str
    chunk_index: int
    section: str
    subsection: str
    page_number: int
    content: str
    embedding: list[float]


@dataclass
class SOPRecord:
    id: str
    sop_id: str
    title: str
    description: str
    department: str
    plant: str
    unit: str
    equipment: str
    category: str
    version: str
    revision_number: int
    status: str
    effective_date: str | None
    review_date: str | None
    expiry_date: str | None
    author: str
    reviewer: str | None
    approver: str | None
    approval_date: str | None
    classification: str
    source_type: str
    source_location: str | None
    file_object_key: str
    file_name: str
    mime_type: str
    file_size: int
    checksum: str
    created_by: str
    updated_by: str
    created_at: str
    updated_at: str
    chunks: list[SOPChunkRecord] = field(default_factory=list)

    def public(self) -> dict[str, Any]:
        result = asdict(self)
        result.pop("chunks", None)
        result["chunk_count"] = len(self.chunks)
        return result


class SOPService:
    def __init__(self, storage: ObjectStorage | None = None) -> None:
        self.storage = storage or ObjectStorage()
        self.records: dict[str, SOPRecord] = {}
        self.audit: dict[str, list[dict[str, Any]]] = {}
        self._load_local_records()

    def _load_local_records(self) -> None:
        if not SessionLocal or not isinstance(self.storage, ObjectStorage):
            return
        with SessionLocal() as session:
            for item in session.query(SOP).all():
                chunks = [
                    SOPChunkRecord(
                        id=chunk.id, sop_id=item.id, sop_version=chunk.sop_version,
                        chunk_index=chunk.chunk_index, section=chunk.section or "Unclassified",
                        subsection=chunk.subsection or "", page_number=chunk.page_number or 1,
                        content=chunk.content, embedding=[],
                    )
                    for chunk in session.query(SOPChunk).filter(SOPChunk.sop_id == item.id).all()
                ]
                self.records[item.id] = SOPRecord(
                    id=item.id, sop_id=item.sop_id, title=item.title, description=item.description or "",
                    department=item.department or "", plant=item.plant or "", unit=item.unit or "",
                    equipment=item.equipment or "", category=item.category or "", version=item.version,
                    revision_number=item.revision_number, status=item.status,
                    effective_date=item.effective_date.isoformat() if item.effective_date else None,
                    review_date=item.review_date.isoformat() if item.review_date else None,
                    expiry_date=item.expiry_date.isoformat() if item.expiry_date else None,
                    author=item.author or "unknown", reviewer=item.reviewer, approver=item.approver,
                    approval_date=item.approval_date.isoformat() if item.approval_date else None,
                    classification=item.classification, source_type=item.source_type,
                    source_location=item.source_location, file_object_key=item.file_object_key,
                    file_name=item.file_name, mime_type=item.mime_type, file_size=item.file_size,
                    checksum=item.checksum, created_by="unknown", updated_by="unknown",
                    created_at=item.created_at.isoformat(), updated_at=item.updated_at.isoformat(),
                    chunks=chunks,
                )

    def _persist_record(self, record: SOPRecord) -> None:
        if not SessionLocal or not isinstance(self.storage, ObjectStorage):
            return
        with SessionLocal() as session:
            item = session.query(SOP).filter(SOP.id == record.id).first()
            if not item:
                item = SOP(id=record.id, sop_id=record.sop_id, title=record.title)
                session.add(item)
            for field_name in (
                "description", "department", "plant", "unit", "equipment", "category", "version",
                "revision_number", "status", "classification", "source_type", "source_location",
                "file_object_key", "file_name", "mime_type", "file_size", "checksum",
            ):
                setattr(item, field_name, getattr(record, field_name))
            item.author = record.author
            item.approver = record.approver
            item.approval_date = datetime.fromisoformat(record.approval_date) if record.approval_date else None
            session.query(SOPChunk).filter(SOPChunk.sop_id == record.id).delete()
            for chunk in record.chunks:
                session.add(SOPChunk(
                    id=chunk.id, sop_id=record.id, sop_version=record.version,
                    chunk_index=chunk.chunk_index, section=chunk.section,
                    subsection=chunk.subsection, page_number=chunk.page_number,
                    content=chunk.content, embedding=None,
                ))
            session.commit()

    def _persist_sop_metadata(self, record: SOPRecord, user: dict[str, Any]) -> None:
        if not SessionLocal or not isinstance(self.storage, ObjectStorage):
            return
        with SessionLocal() as session:
            file_path = str(self.storage.root / record.file_object_key)
            document = session.query(Document).filter(Document.document_id == record.id).first()
            if not document:
                session.add(Document(
                    document_id=record.id, filename=record.file_name,
                    original_filename=record.file_name, file_path=file_path,
                    document_type="sop", object_key=record.file_object_key,
                    mime_type=record.mime_type, size_bytes=record.file_size,
                    uploaded_by=user.get("sub"), status=record.status, version=record.version,
                    sha256=record.checksum,
                ))
            session.add(SOPDocument(
                document_id=record.id, sop_number=record.sop_id, title=record.title,
                description=record.description, department=record.department,
                version=record.version, approval_status=record.status,
            ))
            session.commit()

    def create(
        self,
        payload: bytes,
        filename: str,
        mime_type: str,
        metadata: dict[str, Any],
        user: dict[str, Any],
    ) -> SOPRecord:
        sop_id = str(metadata["sop_id"]).strip()
        version = str(metadata["version"]).strip()
        if not sop_id or not version:
            raise ValueError("SOP ID and version are required.")
        if any(item.sop_id == sop_id and item.version == version for item in self.records.values()):
            raise ValueError("This SOP ID and version already exists.")
        suffix = Path(filename).suffix.lower()
        if suffix not in {".pdf", ".docx"}:
            raise ValueError("SOP files must be PDF or DOCX.")
        if not payload:
            raise ValueError("The SOP file is empty.")
        record_id = str(uuid.uuid4())
        safe_name = Path(filename).name
        key = f"sops/{sop_id}/{version}/{safe_name}"
        stored = self.storage.put(key, payload, mime_type or "application/octet-stream")
        temp_path = self.storage.root / ".processing" / record_id / safe_name
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path.write_bytes(payload)
        try:
            pages = self._extract(temp_path, safe_name)
        finally:
            temp_path.unlink(missing_ok=True)
        now = datetime.now(timezone.utc).isoformat()
        record = SOPRecord(
            id=record_id, sop_id=sop_id, title=str(metadata.get("title") or safe_name),
            description=str(metadata.get("description") or ""),
            department=str(metadata.get("department") or ""), plant=str(metadata.get("plant") or ""),
            unit=str(metadata.get("unit") or ""), equipment=str(metadata.get("equipment") or ""),
            category=str(metadata.get("category") or ""), version=version,
            revision_number=int(metadata.get("revision_number") or 1), status="UPLOADED",
            effective_date=metadata.get("effective_date"), review_date=metadata.get("review_date"),
            expiry_date=metadata.get("expiry_date"), author=str(metadata.get("author") or user.get("sub", "unknown")),
            reviewer=None, approver=None, approval_date=None,
            classification=str(metadata.get("classification") or "INTERNAL"),
            source_type="LOCAL_UPLOAD", source_location=metadata.get("source_location"),
            file_object_key=stored["object_key"], file_name=safe_name,
            mime_type=mime_type or "application/octet-stream", file_size=stored["size_bytes"],
            checksum=stored["sha256"], created_by=str(user.get("sub", "unknown")),
            updated_by=str(user.get("sub", "unknown")), created_at=now, updated_at=now,
        )
        record.chunks = self._chunks(record, pages)
        self.records[record.id] = record
        self._persist_record(record)
        self._persist_sop_metadata(record, user)
        self._event(record, "SOP_UPLOAD", user, {"version": version})
        return record

    def list(self, user: dict[str, Any], filters: dict[str, str | None]) -> list[dict[str, Any]]:
        records = [record for record in self.records.values() if self.authorized(record, user)]
        for key, value in filters.items():
            if value:
                records = [record for record in records if value.lower() in str(getattr(record, key, "")).lower()]
        return [record.public() for record in sorted(records, key=lambda item: item.updated_at, reverse=True)]

    def get(self, record_id: str, user: dict[str, Any]) -> SOPRecord:
        record = self.records.get(record_id)
        if not record or not self.authorized(record, user):
            raise KeyError("SOP not found.")
        self._event(record, "SOP_VIEW", user)
        return record

    def transition(self, record_id: str, action: str, user: dict[str, Any]) -> SOPRecord:
        record = self.get(record_id, user)
        roles = set(user.get("roles", []))
        if action not in TRANSITIONS or not roles.intersection(ROLE_ACTIONS[action]):
            raise PermissionError("You are not authorized to perform this SOP workflow action.")
        expected, target = TRANSITIONS[action]
        if record.status != expected:
            raise ValueError(f"SOP must be {expected} before it can be {action}.")
        if action == "activate":
            for other in self.records.values():
                if other.id != record.id and other.sop_id == record.sop_id and other.status == "ACTIVE":
                    other.status = "SUPERSEDED"
                    other.updated_at = datetime.now(timezone.utc).isoformat()
                    self._event(other, "SOP_SUPERSEDED", user, {"replacement": record.id})
        record.status = target
        record.updated_by = str(user.get("sub", "unknown"))
        record.updated_at = datetime.now(timezone.utc).isoformat()
        if action == "approve":
            record.approver = str(user.get("sub", "unknown"))
            record.approval_date = record.updated_at
        self._event(record, f"SOP_{target}", user)
        self._persist_record(record)
        return record

    def versions(self, record_id: str, user: dict[str, Any]) -> list[dict[str, Any]]:
        record = self.get(record_id, user)
        return [item.public() for item in sorted(self.records.values(), key=lambda item: item.version)
                if item.sop_id == record.sop_id and self.authorized(item, user)]

    def search(self, query: str, user: dict[str, Any], include_historical: bool = False) -> list[dict[str, Any]]:
        records = [record for record in self.records.values() if self.authorized(record, user)
                   and (include_historical or record.status in {"APPROVED", "ACTIVE"})]
        terms = set(re.findall(r"[a-z0-9]+", query.lower()))
        ranked: list[tuple[int, SOPRecord, SOPChunkRecord | None]] = []
        for record in records:
            for chunk in record.chunks:
                haystack = f"{record.title} {record.sop_id} {record.department} {record.plant} {record.unit} {chunk.content}".lower()
                score = sum(1 for term in terms if term in haystack)
                if score:
                    ranked.append((score, record, chunk))
        ranked.sort(key=lambda item: item[0], reverse=True)
        return [{"sop": record.public(), "score": score, "chunk": asdict(chunk) if chunk else None}
                for score, record, chunk in ranked[:10]]

    def audit_for(self, record_id: str, user: dict[str, Any]) -> list[dict[str, Any]]:
        self.get(record_id, user)
        return self.audit.get(record_id, [])

    def document(self, record_id: str, user: dict[str, Any]) -> tuple[SOPRecord, bytes]:
        record = self.get(record_id, user)
        return record, self.storage.get(record.file_object_key)

    def authorized(self, record: SOPRecord, user: dict[str, Any]) -> bool:
        roles = set(user.get("roles", []))
        return bool({"ADMIN", "AUDITOR"} & roles) or record.created_by in {None, user.get("sub")}

    def _event(self, record: SOPRecord, action: str, user: dict[str, Any], metadata: dict[str, Any] | None = None) -> None:
        self.audit.setdefault(record.id, []).append({
            "timestamp": datetime.now(timezone.utc).isoformat(), "user": user.get("sub"),
            "action": action, "resource": "SOP", "resource_id": record.id,
            "success": True, "metadata": metadata or {},
        })

    @staticmethod
    def _extract(path: Path, filename: str) -> list[tuple[int, str]]:
        if path.suffix.lower() == ".docx":
            from docx import Document
            return [(1, "\n".join(paragraph.text for paragraph in Document(path).paragraphs))]
        structured = extract_document(path, demo_mode=False)
        return [(page.page, page.text) for page in structured.pages]

    @staticmethod
    def _chunks(record: SOPRecord, pages: list[tuple[int, str]]) -> list[SOPChunkRecord]:
        chunks: list[SOPChunkRecord] = []
        for page_number, text in pages:
            parts = [part.strip() for part in re.split(r"\n{2,}|(?<=[.!?])\s+(?=[A-Z])", text) if part.strip()]
            for content in parts:
                chunks.append(SOPChunkRecord(
                    id=str(uuid.uuid4()), sop_id=record.id, sop_version=record.version,
                    chunk_index=len(chunks), section="Unclassified", subsection="",
                    page_number=page_number, content=content, embedding=embed(content),
                ))
        return chunks
